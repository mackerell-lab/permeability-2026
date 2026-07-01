from analysis_utlis import *
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_levene_results_to_excel(
    memb_csv_files,
    output_excel="Levene_variance_results.xlsx",
    metric_ls_full=None,
    float_format="%.3g",
    sheet_name="Levene_variance"
):
    if metric_ls_full is None:
        metric_ls_full = ['R2', 'RMSE', 'MAE', 'PearsonR', 'SpearmanRho']

    table_gap = 1          # columns between tables
    start_row = 1
    start_col_left = 1

    writer = pd.ExcelWriter(output_excel, engine="openpyxl", mode="w")
    ws_name = sheet_name

    # Track where to place tables
    table_idx = 0

    for csv_file in memb_csv_files:
        df = concat_metric_files(csv_file)
        levene_df = variance_analysis_by_split(df, metric_ls_full)

        title = Path(csv_file[0]).stem
        n_cols = levene_df.shape[1]
        table_width = n_cols

        # --- Determine left/right placement ---
        if table_idx % 2 == 0:
            col_start = start_col_left
            row_start = start_row
        else:
            col_start = start_col_left + table_width + table_gap
            row_start = start_row

        # --- Title row ---
        pd.DataFrame([[title]]).to_excel(
            writer,
            sheet_name=ws_name,
            startrow=row_start - 1,
            startcol=col_start - 1,
            index=False,
            header=False
        )

        # --- Table ---
        levene_df.to_excel(
            writer,
            sheet_name=ws_name,
            startrow=row_start,
            startcol=col_start - 1,
            index=False,
            float_format=float_format
        )

        # --- Move to next row after 2 tables ---
        if table_idx % 2 == 1:
            start_row += len(levene_df) + 4

        table_idx += 1

    writer.close()

    # =========================
    # Formatting with openpyxl
    # =========================
    wb = load_workbook(output_excel)
    ws = wb[ws_name]

    center = Alignment(horizontal="center", vertical="center")
    bold = Font(bold=True)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    table_idx = 0
    row_cursor = 1

    for csv_file in memb_csv_files:
        df = concat_metric_files(csv_file)
        levene_df = variance_analysis_by_split(df, metric_ls_full)

        n_rows = len(levene_df) + 1  # header included
        n_cols = levene_df.shape[1]

        if table_idx % 2 == 0:
            col_start = start_col_left
        else:
            col_start = start_col_left + n_cols + table_gap

        # --- Merge title ---
        ws.merge_cells(
            start_row=row_cursor,
            start_column=col_start,
            end_row=row_cursor,
            end_column=col_start + n_cols - 1
        )

        title_cell = ws.cell(row=row_cursor, column=col_start)
        title_cell.font = bold
        title_cell.alignment = center

        # --- Apply borders & alignment ---
        for r in range(row_cursor + 1, row_cursor + n_rows + 1):
            for c in range(col_start, col_start + n_cols):
                cell = ws.cell(row=r, column=c)
                cell.alignment = center
                cell.border = border

        # --- Title border ---
        for c in range(col_start, col_start + n_cols):
            cell = ws.cell(row=row_cursor, column=c)
            cell.border = border

        if table_idx % 2 == 1:
            row_cursor += n_rows + 3

        table_idx += 1

    # --- Auto column width ---
    for col in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 2

    wb.save(output_excel)
